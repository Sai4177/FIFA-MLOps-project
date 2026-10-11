"""Builds the single Unit-II submission document (Word) from the real project files/results.
Usage: python docs/build_assignment_doc.py  ->  docs/MLOps_Unit_II_Assignment.docx"""
import json
import subprocess
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "MLOps_Unit_II_Assignment.docx"
doc = Document()

sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = sec.bottom_margin = Inches(0.8)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)
for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
    st = doc.styles[f"Heading {lvl}"]
    st.font.name = "Calibri"; st.font.size = Pt(size); st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def para(text, bold=False, italic=False, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet").paragraph_format.space_after = Pt(2)


def code(text, size=7.5):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; shade(c, "F3F4F6")
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(text.rstrip("\n").split("\n")):
        r = p.add_run(line if i == 0 else "\n" + line)
        r.font.name = "Consolas"; r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table(df, widths=None, font=8.5, header_fill="1F3A5F"):
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, col in enumerate(df.columns):
        c = t.rows[0].cells[i]; c.text = ""; shade(c, header_fill)
        r = c.paragraphs[0].add_run(str(col)); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(255, 255, 255)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def screenshot_slot(n, title, how):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    trPr = t.rows[0]._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    c = t.rows[0].cells[0]; shade(c, "FFF8E1"); c.text = ""
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"\n[ INSERT SCREENSHOT {n} HERE ]\n"); r.bold = True; r.font.size = Pt(11)
    p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"{title}\n{how}\n\n\n\n"); r2.font.size = Pt(9); r2.italic = True
    cap = para(f"Figure {n}. {title}", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


# ------------------------------------------------------------------ data from real runs
versions = json.loads((ROOT / "data" / "versions.json").read_text())["datasets"]
cmp = pd.read_csv(ROOT / "reports" / "preview_model_comparison.csv")
ver = pd.read_csv(ROOT / "reports" / "preview_registry_versions.csv")
quality = json.loads((ROOT / "reports" / "data_quality.json").read_text())
raw = pd.read_csv(ROOT / "data" / "raw" / "fifa_ranking_2022-10-06.csv")
proc = pd.read_csv(ROOT / "data" / "processed" / "fifa_processed_v2.csv")
feat = pd.read_csv(ROOT / "data" / "features" / "fifa_features_v3.csv")
best = cmp.iloc[0]["model"]

# ------------------------------------------------------------------ title
para("MLOps – Unit II Assignment", bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("End-to-End MLOps Workflow: Data Versions → Six Models → MLflow Tracking → Model Registry",
     size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Dataset: FIFA Men's World Ranking (2022-10-06) — fifa_ranking_2022-10-06.csv", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
table(pd.DataFrame([["Student name", "______________________"], ["Roll / Reg. No.", "______________________"],
                    ["Course / Unit", "MLOps – Unit II"], ["Project folder", "fifa-mlops/"]], columns=["Field", "Value"]),
      widths=[1.8, 4.5], font=10)

doc.add_heading("Contents of this document", 1)
bullets(["1. Problem statement and dataset (Task 1)",
         "2. Step-by-step explanation with commands: data engineering, feature engineering, dataset versions, six models, "
         "MLflow tracking, model comparison, registry Versions 1-3 (Tasks 2-9)",
         "3. Screenshots: MLflow tracking and Model Registry",
         "4. Complete code (appendix)"])
p = para("Note on results: the tables below were produced by running the project's own modelling code "
         f"(scikit-learn {__import__('sklearn').__version__}, random_state=42, stratified 80/20 split). "
         "scripts/exp1_mlflow_registry.py runs the identical training code and logs every run to MLflow, so the MLflow UI "
         "should show the same values. The MLflow screenshots in Section 3 must be captured from your own MLflow UI.",
         italic=True, size=9.5)

# ------------------------------------------------------------------ 1 problem
doc.add_heading("1. Problem statement and dataset (Task 1)", 1)
para("The raw dataset lists 211 national football teams from the FIFA ranking published on 6 Oct 2022, with their current and "
     "previous rank and points and their confederation (association).")
table(pd.DataFrame([[c, str(raw[c].dtype), str(raw[c].iloc[0])] for c in raw.columns],
                   columns=["Column", "Type", "Example (Brazil)"]), widths=[1.8, 1.2, 2.5])
para("ML problem (multi-class classification): predict a team's current FIFA tier — Elite (rank 1–20), Strong (21–60), "
     "Mid (61–120) or Developing (121+) — using only the previous ranking snapshot: previous_rank, previous_points and association. "
     "The current rank and points are used only to create the label, never as features, so the model does not leak the answer. "
     "Class counts: " + ", ".join(f"{k} = {v}" for k, v in quality["tier_counts"].items()) + ".")

# ------------------------------------------------------------------ 2 steps
doc.add_heading("2. Step-by-step explanation (Tasks 2–9)", 1)

doc.add_heading("Step 1 – Data engineering and preprocessing (Task 2)", 2)
bullets([
    "Schema validation: all 7 expected columns must exist.",
    "Text cleaning: strip spaces, Unicode-normalise names (Türkiye, Côte d'Ivoire, Curaçao), upper-case codes.",
    "Type enforcement: ranks → integers, points → floats; missing values and duplicate teams removed "
    f"(rows in {quality['rows_in']}, dropped missing {quality['rows_dropped_missing']}, dropped duplicates {quality['rows_dropped_duplicates']}, rows out {quality['rows_out']}).",
    "Business-rule checks: known associations only, points > 0, ranks ≥ 1, ranks unique; "
    f"{quality['points_iqr_outliers']} IQR outliers found in points (kept – values are legitimate).",
    "Label creation: tier = rank_to_tier(rank) using tier_logic.py. Analysis-only columns rank_change and points_change are added but never used as features.",
])
code("python scripts/prepare_data.py          # runs data engineering + feature engineering\n"
     "# -> data/processed/fifa_processed_v2.csv , reports/data_quality.json")

doc.add_heading("Step 2 – Feature engineering (Task 3)", 2)
para("A scikit-learn transformer (FifaFeatureEngineer) derives these features from previous_rank, previous_points and association:")
bullets(["log_prev_points, log_prev_rank – reduce skew", "pct_of_leader – previous points ÷ points of the leader",
         "confed_mean_prev_points, points_vs_confed_mean, points_z_in_confed, confed_size – strength relative to the team's confederation"])
para("Numeric features are standardised and association is one-hot encoded inside the pipeline. Because the transformer is part of the "
     "saved model, the registered model accepts raw inputs and applies exactly the same feature logic at serving time.")

doc.add_heading("Step 3 – Dataset versioning (Task 4)", 2)
para("Each stage is written to its own file; data/versions.json stores row/column counts, parent version and a SHA-256 hash, and every MLflow run is tagged with the three dataset versions and logs the three datasets as inputs.")
table(pd.DataFrame([[v["version"], v["stage"], v["file"], f'{v["rows"]} × {v["columns"]}', v["parent"] or "–", v["sha256"][:16] + "…"] for v in versions],
                   columns=["Version", "Stage", "File", "Rows × Cols", "Parent", "SHA-256"]), widths=[0.6, 1.2, 2.3, 0.8, 0.6, 1.3], font=8)

doc.add_heading("Step 4 – Train six different models (Task 5)", 2)
para("Stratified 80/20 split (168 training / 43 test rows, random_state = 42). Every model sits in the same pipeline, so preprocessing is identical:")
table(pd.DataFrame([["1", "Logistic Regression", "C=1.0, max_iter=2000"], ["2", "K-Nearest Neighbours", "n_neighbors=5"],
                    ["3", "Support Vector Machine", "RBF kernel, C=1.0"], ["4", "Decision Tree", "max_depth=5"],
                    ["5", "Random Forest", "n_estimators=100"], ["6", "Gradient Boosting", "n_estimators=100, learning_rate=0.1"]],
                   columns=["#", "Model", "Version-1 hyper-parameters"]), widths=[0.4, 2.2, 3.2])

doc.add_heading("Step 5 – Track everything with MLflow (Task 6)", 2)
para("For every model one MLflow run logs: parameters (all classifier hyper-parameters, split sizes, seed), metrics "
     "(accuracy, macro precision, macro recall, macro F1, train accuracy, 5-fold CV accuracy mean/std), artifacts "
     "(confusion-matrix PNG/CSV, dataset_versions.json), tags (model family, dataset versions) and the complete model pipeline with signature and input example.")
code("python scripts/exp1_mlflow_registry.py\n"
     "mlflow ui --backend-store-uri sqlite:///mlflow.db --port 5000     # http://127.0.0.1:5000")

doc.add_heading("Step 6 – Compare the models (Task 7)", 2)
para("Selection rule: highest held-out macro-F1, ties broken by 5-fold cross-validated accuracy.")
disp = cmp[["rank", "model", "accuracy", "precision_macro", "recall_macro", "f1_macro", "cv_accuracy_mean", "cv_accuracy_std"]].copy()
disp.columns = ["Rank", "Model", "Accuracy", "Precision", "Recall", "F1 (macro)", "CV acc.", "CV std"]
table(disp, widths=[0.4, 1.55, 0.7, 0.7, 0.65, 0.75, 0.65, 0.55], font=8)
doc.add_picture(str(ROOT / "reports" / "preview_model_comparison.png"), width=Inches(5.6))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para("Figure A. Test accuracy and macro-F1 of the six models.", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
para(f"Result: {best} is the best-performing model (accuracy {cmp.iloc[0]['accuracy']:.2%}, macro-F1 {cmp.iloc[0]['f1_macro']:.4f}); "
     "Gradient Boosting reaches the same test score but a lower cross-validated accuracy. Tree-based models clearly beat "
     "KNN and SVM, whose distance/kernel geometry suits the sharp rank-boundary classes less well. The test set has only 43 teams, "
     "so one team ≈ 2.3 percentage points; the CV columns give a more stable second opinion (Decision Tree has the highest CV accuracy).",
     bold=False)

doc.add_heading("Step 7 – Register the best model as Version 1 (Task 8)", 2)
code('mlflow.register_model(best["model_uri"], "FIFA_Tier_Classifier")   # -> Version 1\n'
     '# alias used by the REST API:  models:/FIFA_Tier_Classifier@champion')

doc.add_heading("Step 8 – Change hyper-parameters, retrain, register Version 2, 3 (Task 9)", 2)
para(f"The best model ({best}) is retrained with two new hyper-parameter sets; each run is logged and registered as a new version of the same registered model. "
     "The alias champion is moved to the best version (ties → newest).")
v = ver.copy(); v.columns = ["Registry version", "Label", "Hyper-parameters", "Accuracy", "F1 (macro)", "CV acc."]
table(v, widths=[0.8, 0.7, 2.9, 0.7, 0.7, 0.6], font=8)

# ------------------------------------------------------------------ 3 screenshots
doc.add_heading("3. Screenshots (MLflow tracking and Model Registry)", 1)
para("Capture these from your MLflow UI after running scripts/exp1_mlflow_registry.py and replace each yellow box (Insert → Picture).", italic=True)
slots = [
    ("Terminal output of prepare_data.py (dataset versions)", "Run: python scripts/prepare_data.py"),
    ("MLflow experiment FIFA_Tier_Classification – runs table", "http://127.0.0.1:5000 → Experiments → FIFA_Tier_Classification (show all 9 runs)"),
    ("MLflow – comparison of the six models", "Select the 6 'comparison' runs → Compare → metrics / parallel-coordinates view"),
    ("MLflow – one run: parameters, metrics and tags", "Open the RandomForest run → Overview"),
    ("MLflow – run artifacts (model, confusion matrix, dataset_versions.json)", "Open the run → Artifacts tab"),
    ("MLflow Model Registry – FIFA_Tier_Classifier with Versions 1, 2, 3", "Models → FIFA_Tier_Classifier"),
    ("MLflow Model Registry – champion alias and version details", "Open a version → show alias, source run, description tag"),
]
for i, (t, h) in enumerate(slots, start=1):
    screenshot_slot(i, t, h)

# ------------------------------------------------------------------ appendix
doc.add_page_break()
doc.add_heading("4. Complete code (appendix)", 1)
files = ["tier_logic.py", "src/fifa_mlops/config.py", "src/fifa_mlops/versioning.py", "src/fifa_mlops/data_engineering.py",
         "src/fifa_mlops/feature_engineering.py", "src/fifa_mlops/modeling.py", "src/fifa_mlops/tracking.py",
         "scripts/prepare_data.py", "scripts/exp1_mlflow_registry.py"]
for f in files:
    doc.add_heading(f, 3)
    code((ROOT / f).read_text(), size=7)

doc.save(OUT)
print("saved", OUT)
